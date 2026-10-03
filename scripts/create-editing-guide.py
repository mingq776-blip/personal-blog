from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

out = Path(r"E:/codex/个人博客/docs/网页内容编辑指南.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 28, "145A2A"), ("Heading 1", 19, "237A3A"), ("Heading 2", 14, "4D8B5F")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

def numbers(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")

def note(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("237A3A")

doc.add_paragraph("网页内容编辑指南", style="Title")
doc.add_paragraph("项目位置：E:\\codex\\个人博客\n线上地址：https://personal-blog-gamma-six.vercel.app")

doc.add_heading("一、最容易修改的内容", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Light Shading Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "想修改的内容"
hdr[1].text = "文件位置"
hdr[2].text = "说明"
for row in [
    ("网站名称、作者、角色、邮箱", "src\\data\\site.ts", "统一修改，全站会同步更新"),
    ("顶部导航栏目", "src\\data\\site.ts", "navigation 数组中修改"),
    ("个人介绍文字", "src\\pages\\about.astro", "直接修改中文段落"),
    ("首页标题和栏目文案", "src\\pages\\index.astro", "首页所有分区的文字"),
    ("项目经历", "src\\data\\experience.ts", "时间、职责、描述和标签"),
    ("作品展示", "src\\content\\projects", "每个 .md 文件是一个作品"),
    ("数据图表", "src\\data\\metrics.ts", "图表名称、数值、颜色和标签"),
    ("文章内容", "src\\content\\posts", "每个 .md 文件是一篇文章"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = row

doc.add_heading("二、推荐方式：直接让 AI 修改", level=1)
doc.add_paragraph("最简单的方法是不自己找文件，直接告诉 Codex 要改什么。例如：")
bullets([
    "“把网站里的‘我的名字’全部改成‘李明’。”",
    "“把个人介绍改成：我是……，目前关注……”",
    "“把项目经历替换成我下面提供的三段经历。”",
    "“把作品展示里的示例项目改成我的三个真实项目。”",
    "“把数据图表里的月均学习时间改成 60 小时。”",
    "“把首页主标题改成‘用作品记录成长’。”",
])
note("你只需要提供实际内容和要求，AI 可以完成查找文件、修改、构建和上传。")

doc.add_heading("三、手动编辑流程", level=1)
numbers([
    "打开 E:\\codex\\个人博客 目录。",
    "用 VS Code 或 Codex 打开上面表格对应的文件。",
    "只修改中文文字、数字、链接或数组内容，不要删除引号、逗号或括号。",
    "双击 build-site.cmd 检查网页是否能正常构建。",
    "告诉 Codex 或运行 git push 上传，Vercel 会自动发布。",
])

doc.add_heading("四、作品或文章怎么添加", level=1)
doc.add_paragraph("作品展示：在 src\\content\\projects 新建 .md 文件。")
doc.add_paragraph("文章内容：在 src\\content\\posts 新建 .md 文件。")
doc.add_paragraph("文件最上方需要填写元数据，例如：")
p = doc.add_paragraph()
r = p.add_run('---\ntitle: "作品名称"\nsummary: "一句话介绍"\nyear: 2026\nrole: "设计与开发"\nstack: ["Astro", "设计"]\ncover: "../../assets/covers/project-example.webp"\ncoverAlt: "封面描述"\nfeatured: true\ndemo: false\ndraft: false\n---')
r.font.name = "Consolas"
r.font.size = Pt(9)

doc.add_heading("五、图片怎么替换", level=1)
bullets([
    "作品封面：src\\assets\\covers",
    "绿色风景图片：src\\assets\\landscapes",
    "网站图标：public\\favicon.svg",
    "分享缩略图：public\\images\\og-default.webp",
    "建议图片使用 WebP，横向比例 16:10 或 16:9。",
])
doc.add_heading("六、不要修改的内容", level=1)
bullets([
    "不要修改 node_modules、dist、.astro 或 .tools 目录。",
    "不要删除 package.json、pnpm-lock.yaml 或 astro.config.mjs 中的结构代码。",
    "如果修改后构建报错，立即告诉 Codex 查看错误并恢复。",
])
doc.add_heading("七、发布后如何检查", level=1)
numbers([
    "本地预览：双击 preview-site.cmd。",
    "检查首页、四个导航栏目和作品详情。",
    "GitHub 上传后等待 Vercel 显示 Ready。",
    "中国大陆访问 Vercel 可能需要代理；本地预览不受影响。",
])
doc.save(out)
print(out)
