from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
out = Path(r'E:\codex\个人博客\docs\绿色版网站说明.docx')
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles['Normal']; normal.font.name = 'Microsoft YaHei'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); normal.font.size = Pt(10.5); normal.paragraph_format.line_spacing = 1.4
for name, size, color in [('Title', 28, '145A2A'), ('Heading 1', 19, '237A3A')]:
    s = doc.styles[name]; s.font.name = 'Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color); s.font.bold = True
def bullets(items):
    for item in items: doc.add_paragraph(item, style='List Bullet')
doc.add_paragraph('绿色版网站说明', style='Title')
doc.add_paragraph('版本：2026-10-03\n线上地址：https://personal-blog-gamma-six.vercel.app')
doc.add_heading('一、绿色视觉系统', level=1)
bullets([
    '整体主色改为自然绿色，搭配浅绿背景、深绿文字和柔和光线。',
    '默认使用明亮的浅色主题，不再默认进入深色页面；仍保留深色主题切换。',
    '首页加入绿色粒子网络、绿色标题描边和自然光晕。',
    '数据图表统一改为不同层次的绿色。',
])
doc.add_heading('二、绿色风景图片', level=1)
bullets([
    '首页增加 3 张绿色风景图片：森林、湖泊山谷和绿色草地。',
    '森林图片作为全站固定背景，经过模糊、低透明度和顶部到底部渐隐处理。',
    '首页主视觉后方加入湖水山谷图片，作为粒子动画的绿色背景。',
    '图片使用 WebP 压缩和响应式尺寸，保留文字可读性。',
    '图片来源为 Unsplash，可继续替换为个人拍摄或 AI 生成图片。',
])
doc.add_heading('三、英文与内容调整', level=1)
bullets([
    '保留少量英文标签：GREEN NOTES、SELECTED WORKS、PROJECT MOMENTUM、IMAGE / UNSPLASH。',
    '英文仅作为视觉层次和装饰，不替代中文正文。',
    '主导航保持个人介绍、项目经历、作品展示、数据图表四个核心栏目。',
])
doc.add_heading('四、性能结果', level=1)
bullets([
    '首页 Lighthouse：性能 92、无障碍 100、最佳实践 100、SEO 100。',
    '数据页 Lighthouse：性能 100、无障碍 100、最佳实践 100、SEO 100。',
    '继续使用 Canvas 2D 和 CSS transform，移除 React 与 Three.js。',
])
doc.add_heading('五、设计预览', level=1)
for image, caption in [
    (r'E:\codex\个人博客\docs\preview-home-v4.png', '绿色版首页：绿色主色与自然风景图片'),
    (r'E:\codex\个人博客\docs\preview-dashboard-v4.png', '绿色版数据页：统一绿色图表'),
]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(image, width=Inches(6.2))
    c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.save(out)
print(out)
