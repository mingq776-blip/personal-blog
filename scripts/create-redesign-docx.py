from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out = Path(r'E:\codex\个人博客\docs\网站改版说明.docx')
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles['Normal']; normal.font.name = 'Microsoft YaHei'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); normal.font.size = Pt(10.5); normal.paragraph_format.line_spacing = 1.4
for name, size, color in [('Title', 28, '0B1B2B'), ('Heading 1', 19, '087F72')]:
    s = doc.styles[name]; s.font.name = 'Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color); s.font.bold = True
def bullets(items):
    for item in items: doc.add_paragraph(item, style='List Bullet')
doc.add_paragraph('网站改版说明', style='Title')
doc.add_paragraph('版本：2026-10-03\n线上地址：https://personal-blog-gamma-six.vercel.app')
doc.add_heading('一、视觉与动画升级', level=1)
bullets([
    '从极简排版调整为高对比、强标题和杂志式栅格布局。',
    '首页增加轻量 Canvas 2D 粒子网络、轨道旋转、贴纸标签和跑马灯动画。',
    '栏目卡片增加进入动画、悬停浮动、轻微倾斜和图表生长效果。',
    '动画仅使用 Canvas 2D、CSS transform 和 IntersectionObserver，未引入 React 或 Three.js。',
])
doc.add_heading('二、栏目结构精简', level=1)
bullets([
    '首页：杂志封面、个人简介、经历、作品与数据概览。',
    '个人介绍：个人定位、关注方向和做事原则。',
    '项目经历：按时间线展示项目阶段、职责与成果。',
    '作品展示：项目卡片与项目详情。',
    '数据图表：推进趋势、能力分布和时间分配。',
])
doc.add_heading('三、性能优化', level=1)
bullets([
    '移除了 React、Three.js、React Three Fiber 和 Motion 依赖。',
    '粒子数量按设备和屏幕尺寸限制，设备像素比上限为 1.5。',
    'Canvas 在不可见时停止绘制，页面隐藏时停止动画。',
    '移除了造成滚动高度变化的 content-visibility 设置，CLS 为 0。',
    '首页 Lighthouse：性能 99、无障碍 100、最佳实践 100、SEO 100。',
    '数据页 Lighthouse：性能 100、无障碍 100、最佳实践 100、SEO 100。',
])
doc.add_heading('四、项目位置', level=1)
bullets([
    '本地目录：E:\\codex\\个人博客',
    'GitHub：https://github.com/mingq776-blip/personal-blog',
    'Vercel：https://vercel.com/mingq776-blip/personal-blog',
])
doc.add_heading('五、后续替换', level=1)
bullets([
    '将页面中的“我的名字”替换为真实姓名或昵称。',
    '把示例项目、示例经历和数据替换为真实内容。',
    '替换封面图和头像，再运行 build-site.cmd 并 push。',
])
doc.save(out)
print(out)
