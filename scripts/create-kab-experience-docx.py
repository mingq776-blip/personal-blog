from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out = Path(r"E:/codex/个人博客/docs/社团与KAB经历补充说明.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 26, "145A2A"), ("Heading 1", 18, "237A3A")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
doc.add_paragraph("社团与 KAB 经历补充说明", style="Title")
doc.add_paragraph("来源目录：D:\\KAB+社管\n整理日期：2026-10-04")
doc.add_heading("一、已阅读的主要材料", level=1)
bullets([
    "徐晨洋年度工作报告",
    "个人能力证明",
    "团委社团管理部徐晨洋信息表",
    "博雅书院职业分享会、双创领航讲座、大创结项答辩主持词",
    "社团管理部与二课活动相关策划、通知和申请材料",
])
doc.add_heading("二、新增到网站的经历", level=1)
doc.add_paragraph("1. 团委社团管理部｜工作人员")
bullets([
    "参与社团活动材料检查、值班排班、活动宣传和迷彩青春社团活动跟进。",
    "承担就业分享、双创讲座与大学生创新创业训练计划结项答辩等多类主持任务。",
    "网站展示口径：10+ 场校园活动、4 类主持任务、材料检查与排班、活动宣传设计。",
])
doc.add_paragraph("2. KAB 创新创业俱乐部｜综合工作组成员")
bullets([
    "参与就业与创业分享活动策划、嘉宾对接、现场主持和材料整理。",
    "协同推进双创讲座、优秀学生经验分享与互动答疑。",
    "网站展示口径：综合工作组成员、就业分享会主持、4 位分享嘉宾、双创讲座协同。",
])
doc.add_heading("三、经历说明", level=1)
bullets([
    "年度工作报告中的具体工作包括安排值班表、盲盒工作人员、跟进迷彩青春社团活动、制作厨艺大赛海报、检查社团文书和征兵微视频核对。",
    "个人能力证明中记录了 KAB 综合工作组成员身份和优秀创业就业学长经验分享主持经历。",
    "大创结项答辩主持覆盖 8 个项目团队。",
    "高中学生会副主席经历暂未放入主展示区，避免削弱大学阶段经历的重点。",
])
doc.add_heading("四、隐私与投放建议", level=1)
bullets([
    "未公开联系电话、学号和出生年月。",
    "正式投递时可以用 PDF 版证明、主持稿截图或活动现场图片作为佐证。",
    "涉及活动成效时，优先使用有文件或现场记录支撑的数字。",
])
doc.save(out)
print(out)
