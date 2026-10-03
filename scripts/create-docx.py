from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "个人博客使用与维护说明.docx"
NAVY = "0B1B2B"
TEAL = "087F72"
GOLD = "B96B15"
MUTED = "62717A"

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(2.1)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.35

for style_name, size, color in [("Title", 28, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 14, TEAL), ("Heading 3", 11.5, GOLD)]:
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True

if "Code Block" not in styles:
    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code Block"]
code_style.font.name = "Consolas"
code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
code_style.font.size = Pt(9)
code_style.font.color.rgb = RGBColor.from_string("17324D")
code_style.paragraph_format.left_indent = Cm(0.4)
code_style.paragraph_format.space_before = Pt(3)
code_style.paragraph_format.space_after = Pt(8)


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_code(text):
    paragraph = doc.add_paragraph(style="Code Block")
    paragraph.add_run(text)
    shade(paragraph, "F2F5F5")
    return paragraph


def add_bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_caption(text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)


title = doc.add_paragraph(style="Title")
title.add_run("个人博客使用与维护说明")
subtitle = doc.add_paragraph()
run = subtitle.add_run("Astro + React + Three.js 博客与作品集")
run.font.size = Pt(13)
run.font.color.rgb = RGBColor.from_string(TEAL)
doc.add_paragraph("项目位置：E:\\codex\\个人博客\n生成日期：2026-10-03\nGit 分支：main")

info = doc.add_table(rows=1, cols=2)
info.style = "Light Shading Accent 1"
info.rows[0].cells[0].text = "检查项目"
info.rows[0].cells[1].text = "结果"
for label, value in [
    ("Astro 类型检查", "0 errors / 0 warnings"),
    ("静态构建", "26 个页面构建成功"),
    ("搜索索引", "6 个文章与项目详情页，689 个词"),
    ("内部链接检查", "28 个链接，0 个失效"),
    ("首页 Lighthouse", "性能 93 / 无障碍 100 / 最佳实践 100 / SEO 100"),
    ("文章页 Lighthouse", "性能 100 / 无障碍 100 / 最佳实践 100 / SEO 100"),
]:
    cells = info.add_row().cells
    cells[0].text = label
    cells[1].text = value

doc.add_page_break()
doc.add_heading("1. 网站包含什么", level=1)
add_bullets([
    "首页：粒子几何体 3D 主视觉、精选文章、精选项目和经历时间线。",
    "博客：文章列表、分类、标签、归档、文章目录、阅读进度和相关阅读。",
    "项目：项目列表与详情，支持技术栈、年份、职责、访问链接和代码链接。",
    "关于：个人介绍、能力方向和近况，可集中替换为真实资料。",
    "搜索：Pagefind 在构建时生成静态索引，无需服务器或数据库。",
    "主题：暗色优先，可切换浅色主题；支持减少动画与无 WebGL 降级。",
    "SEO：RSS、站点地图、Canonical、Open Graph 和结构化数据。",
])

doc.add_heading("2. 最快的运行方式", level=1)
doc.add_paragraph("项目已经内置便携版 Node.js，不需要修改系统环境。直接双击以下文件：")
add_bullets([
    "start-dev.cmd：安装依赖并启动本地开发服务器。",
    "build-site.cmd：构建静态网站和 Pagefind 搜索索引。",
    "preview-site.cmd：预览已经构建的静态网站。",
])
doc.add_paragraph("开发服务器默认地址通常为：")
add_code("http://localhost:4321/")

doc.add_heading("3. 目录结构", level=1)
add_code("E:\\codex\\个人博客\n├─ src\\content\\posts       文章 Markdown\n├─ src\\content\\projects    项目 Markdown\n├─ src\\assets\\covers       文章和项目封面\n├─ src\\components          页面组件与 3D 组件\n├─ src\\pages               网站路由\n├─ src\\data\\site.ts        网站名称、作者和社交链接\n├─ public                  站点图标、分享图\n├─ docs                    Word 文档、预览图和测试结果\n└─ dist                    构建产物，可部署")

doc.add_page_break()
doc.add_heading("4. 发布一篇新文章", level=1)
add_numbered([
    "在 src\\content\\posts 新建 Markdown 文件，文件名使用英文短横线，例如 my-first-post.md。",
    "复制下面的 Frontmatter，并修改标题、摘要、日期、分类、标签和封面路径。",
    "正文使用标准 Markdown 编写，支持标题、列表、引用、链接、图片和代码块。",
    "运行 build-site.cmd，确认没有错误。",
    "提交并推送 Git，Vercel 会自动重新构建和部署。",
])
add_code("---\ntitle: \"文章标题\"\ndescription: \"文章摘要\"\npubDate: 2026-10-03\nupdatedDate: 2026-10-05\ncategory: \"技术\"\ntags: [\"Astro\", \"写作\"]\ncover: \"../../assets/covers/post-why-blog.webp\"\ncoverAlt: \"封面描述\"\nfeatured: true\ndraft: false\n---\n\n从这里开始写正文。")
doc.add_paragraph("说明：draft 为 true 时不会进入生产页面；featured 为 true 时会优先出现在首页精选区域。")

doc.add_heading("5. 添加或替换项目", level=1)
add_numbered([
    "在 src\\content\\projects 新建 Markdown 文件。",
    "填写项目名称、摘要、年份、职责、技术栈和封面。",
    "把 demo 改为 false，表示它不是示例项目。",
    "有线上地址时填写 liveUrl，有代码仓库时填写 repoUrl。",
    "运行构建并发布。",
])
add_code("---\ntitle: \"项目名称\"\nsummary: \"项目简介\"\nyear: 2026\nrole: \"设计与开发\"\nstack: [\"Astro\", \"React\", \"Three.js\"]\ncover: \"../../assets/covers/project-creative-site.webp\"\ncoverAlt: \"封面描述\"\nliveUrl: \"https://example.com\"\nrepoUrl: \"https://github.com/example/repo\"\nfeatured: true\ndemo: false\ndraft: false\n---")

doc.add_page_break()
doc.add_heading("6. 替换网站身份和链接", level=1)
doc.add_paragraph("统一编辑 src\\data\\site.ts：")
add_bullets([
    "title：网站名称。",
    "author：作者名称。",
    "role：首页身份描述。",
    "description：网站默认摘要。",
    "url：正式部署地址。",
    "navigation：顶部导航。",
    "socials：GitHub、即刻和邮箱等链接。",
])
doc.add_paragraph("同时检查 astro.config.mjs 中的 site 地址，以及 src\\pages\\about.astro 中的个人介绍和近况。")

doc.add_heading("7. 图片和封面", level=1)
add_bullets([
    "文章封面放在 src\\assets\\covers，项目封面也放在同一目录。",
    "建议使用 WebP 或 AVIF，横向比例为 16:10 或 16:9。",
    "Astro 会在构建时生成多种尺寸，减少手机流量并保持清晰度。",
    "替换图片后，确保 Markdown 中的 cover 相对路径正确。",
    "首版封面为原创程序生成的抽象 WebP 占位素材；当前会话未提供内置 AI 生图工具且未配置 API Key，因此没有执行模型生图。",
])

doc.add_heading("8. 构建与本地预览", level=1)
add_code("# 双击 build-site.cmd，或使用项目内便携工具执行\ncorepack pnpm check\ncorepack pnpm build\ncorepack pnpm preview")
add_bullets([
    "check 负责 TypeScript 和 Astro 模板检查。",
    "build 生成静态页面，并在 dist\\pagefind 创建搜索索引。",
    "preview 在本地检查生产构建效果。",
])

doc.add_page_break()
doc.add_heading("9. 部署到 GitHub + Vercel", level=1)
add_numbered([
    "登录 GitHub，创建一个空仓库，例如 personal-blog。",
    "在项目目录配置远程地址：git remote add origin <你的仓库地址>。",
    "推送代码：git push -u origin main。",
    "登录 Vercel，选择 Import Git Repository，导入该仓库。",
    "Framework Preset 选择 Astro；Build Command 使用 pnpm build；Output Directory 使用 dist。",
    "点击 Deploy，等待构建完成，获得免费 *.vercel.app 地址。",
    "把正式地址同步修改到 astro.config.mjs 的 site 和 src\\data\\site.ts 的 url。",
])
add_code("git remote add origin https://github.com/你的用户名/personal-blog.git\ngit push -u origin main")
doc.add_paragraph("如果 Vercel 登录或中国大陆访问测试不稳定，可在 Cloudflare Pages 导入同一仓库。Framework Preset 选择 Astro，构建命令和输出目录相同。")

doc.add_heading("10. 日常发布流程", level=1)
add_numbered([
    "新建或修改 Markdown 文章、项目。",
    "运行 build-site.cmd，确认构建成功。",
    "执行 git add . 和 git commit -m \"post: 新文章标题\"。",
    "执行 git push。",
    "等待 Vercel 自动部署，检查正式地址。",
])

doc.add_heading("11. 备份和迁移", level=1)
add_bullets([
    "文章、项目、页面配置都在 Git 仓库中，GitHub 是第一层云端备份。",
    "可以在任意电脑执行 git clone 恢复完整项目。",
    "dist 是可重新生成的构建结果，不需要作为唯一备份。",
    "迁移到独立域名或国内主机时源码无需重写，只需修改 site、url、DNS 和部署配置。",
    "不要公开提交 .env、密码、API Key 或个人隐私信息。",
])

doc.add_page_break()
doc.add_heading("12. 当前预览", level=1)
for image_name, caption in [
    ("preview-home.png", "首页：3D 粒子主视觉、精选文章、项目和时间线"),
    ("preview-article.png", "文章页：封面、文章目录、阅读进度和正文排版"),
    ("preview-projects.png", "项目页：示例项目卡片与技术栈"),
]:
    image_path = ROOT / "docs" / image_name
    if image_path.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
        add_caption(caption)

doc.add_heading("13. 注意事项", level=1)
add_bullets([
    "示例文章和项目已经明确标记为占位内容，正式上线前应替换或确认。",
    "GitHub 和 Vercel 登录、验证码、域名购买与实名备案需要由本人完成。",
    "免费 Vercel 子域名不备案，中国大陆访问速度可能波动。",
    "首页 3D 会按需加载；如果设备不支持 WebGL，会自动显示静态轨道背景。",
    "中文 Pagefind 搜索不提供词干匹配，但普通关键词检索可以正常使用。",
])

doc.save(OUT)
print(OUT)
