from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out = Path(r"E:/codex/个人博客/docs/国内访问部署方案.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 26, "145A2A"), ("Heading 1", 18, "237A3A")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
def numbers(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")
doc.add_paragraph("国内访问部署方案", style="Title")
doc.add_paragraph("目标：让中国大陆读者更稳定地打开现有网站，并保留 Vercel 版本作为海外镜像。")
doc.add_heading("一、三种方案对比", level=1)
table = doc.add_table(rows=1, cols=5)
table.style = "Light Shading Accent 1"
head = table.rows[0].cells
for i, value in enumerate(["方案", "国内速度", "是否需要备案", "预计成本", "适合情况"]):
    head[i].text = value
for row in [
    ("继续用 Vercel", "不稳定", "不需要", "0 元", "只给面试官或能使用代理的人"),
    ("香港主机 + 独立域名", "中等", "一般不需要", "约 100—300 元/年", "不想备案，但希望比 Vercel 稳定"),
    ("国内对象存储/CDN + 独立域名", "最好", "需要 ICP 备案", "约 100—500 元/年", "长期在国内展示、简历、比赛、答辩"),
]:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
doc.add_heading("二、推荐方案", level=1)
doc.add_paragraph("如果只是保研、面试和作品展示，优先选择“香港主机 + 独立域名”，速度比 Vercel 稳定，通常不需要备案。")
doc.add_paragraph("如果网站需要长期在国内公开访问、放入简历或参加正式项目，选择“国内对象存储/CDN + 独立域名 + ICP 备案”。")
doc.add_heading("三、国内对象存储部署步骤", level=1)
numbers([
    "在阿里云、腾讯云或华为云购买独立域名，并完成实名认证。",
    "提交 ICP 备案，备案主体和网站内容需真实合规。",
    "创建国内 COS/OSS/OBS 存储桶，开启静态网站托管。",
    "把本地 E:\\codex\\个人博客\\dist 目录中的全部文件上传到存储桶。",
    "设置默认首页 index.html，设置 404 文档为 404.html。",
    "绑定独立域名，配置 HTTPS 证书和 CDN 加速。",
    "修改 astro.config.mjs 和 src/data/site.ts 中的正式网址，重新构建并上传。",
    "用 ITDOG、站长工具或 Boce 从多地测试访问速度和 HTTPS 状态。",
])
doc.add_heading("四、香港主机部署步骤", level=1)
numbers([
    "购买香港轻量应用服务器或静态托管服务。",
    "购买独立域名并解析到香港服务器。",
    "使用 Nginx 托管 dist 目录，配置 index.html 和 404.html。",
    "申请免费 HTTPS 证书并开启自动续期。",
    "重新构建网站并上传 dist 文件。",
])
doc.add_heading("五、代码需要改什么", level=1)
bullets([
    "astro.config.mjs：把 site 改成新域名。",
    "src/data/site.ts：把 url 改成新域名。",
    "运行 build-site.cmd，重新生成 dist。",
    "上传 dist 即可，不需要重写网站代码。",
])
doc.add_heading("六、访问地址建议", level=1)
bullets([
    "海外访问：保留 personal-blog-gamma-six.vercel.app。",
    "国内访问：使用自己的独立域名，例如 xuchenyang.com 或 xuchenyang.cn。",
    "简历和 PPT 中优先放独立域名，二维码也指向独立域名。",
])
doc.save(out)
print(out)
