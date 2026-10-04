from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

out = Path(r"E:/codex/个人博客/docs/个人简历与量化项目填充说明.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 28, "145A2A"), ("Heading 1", 19, "237A3A")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

def numbers(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")

doc.add_paragraph("个人简历与量化项目填充说明", style="Title")
doc.add_paragraph("网站：https://personal-blog-gamma-six.vercel.app\n本地项目：E:\\codex\\个人博客\n整理日期：2026-10-04")

doc.add_heading("一、本次使用的内容来源", level=1)
bullets([
    "量化项目来源：可量化项目与技术栈展示清单（2026-10-04）。",
    "个人信息与荣誉来源：徐晨洋个人简历更新版、曾获奖励和自我鉴定。",
    "填充范围：个人介绍、首页、项目经历、作品展示、数据图表和荣誉资格。",
])

doc.add_heading("二、网站已填充内容", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Light Shading Accent 1"
head = table.rows[0].cells
head[0].text, head[1].text, head[2].text = "栏目", "已填充内容", "主要来源"
for row in [
    ("首页", "徐晨洋｜宁夏大学 2025 级大数据管理与应用；让数据变成作品；629 份问卷与 7 个核心项目", "简历 + 量化清单"),
    ("个人介绍", "教育背景、综合测评、排名、奖学金、关注方向、展示原则", "简历更新版"),
    ("项目经历", "数字画笔、中卫烙画、青穗 AI、古建模拟、CMAU、个人博客", "量化清单 + 简历"),
    ("作品展示", "7 个核心项目卡片和详情页，显示量化指标、技术栈、证据口径与边界", "量化清单"),
    ("数据图表", "629 份有效问卷、4 个模型、9 个 API、20000 游戏日、能力证据分布", "量化清单"),
    ("荣誉与资格", "10 项奖学金、竞赛、志愿服务、CET-4 和围棋业余三段", "简历曾获奖励"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = row

doc.add_heading("三、重点展示顺序", level=1)
numbers([
    "青穗 AI 财务教练：4 个模型、9 个 API、3000 条模拟样本。",
    "数字画笔·童心筑梦计划：45 名儿童、135 件作品、出勤率超过 95%。",
    "中卫烙画：263 份有效问卷、α=0.854、KMO=0.816、答辩排名 7/17。",
    "大创古建灾害模拟：2 座古建、20000 游戏日、3 类灾害标签。",
    "个人博客：16+ 页面/路由、11 个核心组件、Lighthouse 最高 100。",
    "Pixel2Motion：IoU=0.9937、102 条 SVG 路径、32 帧 GIF。",
    "CMAU 市场调查：366 份有效问卷、回归与 Bootstrap 检验、校赛二等奖。",
])

doc.add_heading("四、隐私与公开范围", level=1)
bullets([
    "未公开手机号、学号、出生年月和正式邮箱。",
    "公开页面仅展示姓名、学校、专业、项目、荣誉和 GitHub 地址。",
    "后续补充邮箱后，可在 site.ts 中填入，并恢复邮件联系入口。",
])

doc.add_heading("五、数据口径与风险提示", level=1)
bullets([
    "青穗 AI 的所有模型指标来自 3000 条模拟样本，只能写“模拟数据技术验证”。",
    "中卫烙画的预算、盈亏平衡和 12 个月目标属于计划/测算，不是已实现收入。",
    "大创古建结构完整率 100% 属于探索性预实验结果，不替代有限元或真实监测。",
    "CMAU 的 366 份问卷和校赛二等奖来自简历材料，正式投递前应补齐源数据与获奖证明。",
    "团队项目使用“参与/负责相应模块”的表述，正式投递前再次核对实际分工。",
])

doc.add_heading("六、后续建议", level=1)
bullets([
    "补充正式邮箱和个人照片。",
    "整理一份可公开的获奖证书清单或附录。",
    "为青穗接入 30—50 名真实学生小样本测试。",
    "完成一个 SQL + 可视化分析项目，补齐数据分析作品。",
])
doc.save(out)
print(out)
