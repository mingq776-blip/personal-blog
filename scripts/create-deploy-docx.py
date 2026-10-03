from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

out = Path(r'E:\codex\个人博客\docs\GitHub与Vercel上传部署步骤.docx')
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1); section.bottom_margin = Cm(2.1); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles['Normal']; normal.font.name = 'Microsoft YaHei'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); normal.font.size = Pt(10.5); normal.paragraph_format.line_spacing = 1.35
for name, size, color in [('Title', 28, '0B1B2B'), ('Heading 1', 19, '087F72'), ('Heading 2', 14, 'B96B15')]:
    s = doc.styles[name]; s.font.name = 'Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color); s.font.bold = True

def code(text):
    p = doc.add_paragraph(); r = p.add_run(text); r.font.name = 'Consolas'; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string('17324D'); return p

def bullets(items):
    for item in items: doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph('GitHub 与 Vercel 上传部署步骤', style='Title')
doc.add_paragraph('项目：E:\\codex\\个人博客')
doc.add_heading('一、先登录 GitHub', level=1)
doc.add_paragraph('当前已经打开 PowerShell 登录窗口。按窗口提示在浏览器完成 GitHub 授权。如果你还没有 GitHub 账号，先在 github.com 注册。')
doc.add_heading('二、创建 GitHub 仓库', level=1)
doc.add_paragraph('登录完成后运行：')
code('gh repo create personal-blog --public --source . --remote origin --push')
doc.add_paragraph('这条命令会自动创建个人博客仓库、关联远程地址并上传全部代码。')
doc.add_heading('三、手动上传方式', level=1)
code('git remote add origin https://github.com/你的用户名/personal-blog.git\ngit push -u origin main')
doc.add_heading('四、登录 Vercel', level=1)
bullets([
    '打开 vercel.com，选择 Continue with GitHub。',
    '授权 Vercel 访问 GitHub 仓库。',
    '进入 Add New Project，选择 personal-blog。',
    'Framework Preset 选择 Astro。',
    'Build Command 填 pnpm build。',
    'Output Directory 填 dist。',
    '点击 Deploy，等待一两分钟。',
])
doc.add_heading('五、获得网址', level=1)
doc.add_paragraph('部署完成后会得到类似 personal-blog.vercel.app 的免费网址。每次 git push，Vercel 都会自动重新部署。')
doc.add_heading('六、修改正式网址', level=1)
bullets([
    '修改 astro.config.mjs 中的 site。',
    '修改 src/data/site.ts 中的 url。',
    '再次 git add .、git commit、git push。',
])
doc.add_heading('七、注意事项', level=1)
bullets([
    'GitHub 和 Vercel 登录必须由你本人完成。',
    '免费 vercel.app 域名不备案，中国大陆访问速度可能波动。',
    '示例文章和项目上线后可继续替换成真实内容。',
])
doc.save(out)
print(out)
