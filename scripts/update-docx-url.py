from docx import Document
from pathlib import Path

url = 'https://personal-blog-gamma-six.vercel.app'
main = Path(r'E:\codex\个人博客\docs\个人博客使用与维护说明.docx')
deploy = Path(r'E:\codex\个人博客\docs\GitHub与Vercel上传部署步骤.docx')

for path in [main, deploy]:
    doc = Document(path)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith('项目位置：'):
            if '正式地址' not in paragraph.text:
                paragraph.add_run(f'\n正式地址：{url}')
            break
    else:
        doc.add_paragraph(f'正式地址：{url}')
    doc.save(path)
    print(path)
