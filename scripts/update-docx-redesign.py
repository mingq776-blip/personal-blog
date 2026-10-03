from docx import Document
path = r'E:\codex\个人博客\docs\个人博客使用与维护说明.docx'
doc = Document(path)
doc.paragraphs[1].text = 'Astro + Canvas 2D 杂志式个人网站'
items = [
    '首页：杂志式封面、轻量粒子动画、精选项目经历、作品与数据预览。',
    '个人介绍：个人定位、关注方向、创作原则和当前状态。',
    '项目经历：按时间线整理实践阶段、职责、行动和成果。',
    '作品展示：项目卡片、技术栈、项目详情和结果说明。',
    '数据图表：项目推进趋势、能力分布、时间分配和行动结论。',
    '辅助页面：搜索、RSS、Sitemap 和 404 继续保留。',
    '性能优化：移除 React 和 Three.js，改用 Canvas 2D 与 CSS transform，避免滚动布局跳动。',
]
for index, value in zip(range(5, 12), items):
    doc.paragraphs[index].text = value
doc.paragraphs[20].text = doc.paragraphs[20].text.replace('页面组件与 3D 组件', '页面组件与 Canvas 动画')
doc.save(path)
print(path)
