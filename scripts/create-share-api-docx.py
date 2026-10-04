from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
out = Path(r"E:/codex/个人博客/docs/网站分享与API说明.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 26, "145A2A"), ("Heading 1", 18, "237A3A")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
doc.add_paragraph("网站分享与 API 说明", style="Title")
doc.add_heading("一、API 是哪里来的", level=1)
doc.add_paragraph("网站中提到的 9 个 API 来自青穗 AI 财务教练项目自己的 Flask 后端，不是购买的外部接口。")
bullets([
    "源代码位置：E:\\codex\\青穗AI财务教练\\原型\\app.py",
    "后端框架：Python + Flask。",
    "模型计算：scikit-learn 模型负责健康分、异常检测、用户分层和目标预测。",
    "解释能力：Qwen 通过 llm_service 负责对话解释，不负责金融数值计算。",
    "9 个端点：/api/health、/api/plan、/api/score、/api/anomaly、/api/segment、/api/goal、/api/quiz、/api/coach、/api/dashboard。",
])
doc.add_paragraph("如果正式投递，建议写“参与搭建 Flask 后端和 9 个接口原型”，不要写“独立开发全部系统”，除非你能说明每一部分的个人分工。")
doc.add_heading("二、直接分享网址", level=1)
doc.add_paragraph("线上网址：https://personal-blog-gamma-six.vercel.app")
doc.add_paragraph("GitHub：https://github.com/mingq776-blip/personal-blog")
doc.add_paragraph("把网址复制到微信、QQ、邮件、简历或 PPT 中即可分享。")
doc.add_heading("三、扫码分享", level=1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(r"E:/codex/个人博客/docs/网站二维码.png", width=Inches(2.4))
doc.add_paragraph("二维码对应网站地址，可用于 PPT、海报、简历二维码或线下展示。")
doc.add_heading("四、分享时的注意事项", level=1)
bullets([
    "中国大陆网络可能无法稳定直连 Vercel，分享时建议提醒对方使用代理或海外网络。",
    "如果长期在国内展示，建议后续购买独立域名并迁移到香港或国内主机。",
    "不要公开手机号、学号、出生年月等隐私信息。",
    "如果只用于面试或答辩，可以分享线上网址、GitHub 仓库和二维码。",
])
doc.save(out)
print(out)
