from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out = Path(r"E:/codex/个人博客/docs/国内展示版重建方案.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 26, "145A2A"), ("Heading 1", 18, "237A3A")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
def numbers(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")
doc.add_paragraph("国内展示版重建方案", style="Title")
doc.add_heading("一、核心建议", level=1)
doc.add_paragraph("不建议维护两份完全独立的源码。建议保留一套 Astro 代码，同时部署两个版本：")
bullets([
    "海外版：继续部署 Vercel，免费维护。",
    "国内版：部署到腾讯云 COS + CDN，或阿里云 OSS + CDN，并绑定已备案的独立域名。",
    "两版共用文章、项目、图片和数据；只修改部署地址和域名配置。",
])
doc.add_heading("二、推荐技术方案", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Light Shading Accent 1"
head = table.rows[0].cells
for i, value in enumerate(["方案", "优点", "成本", "适合情况"]):
    head[i].text = value
for row in [
    ("腾讯云 COS + CDN", "控制台易用，国内节点稳定", "约 100—300 元/年", "优先推荐"),
    ("阿里云 OSS + CDN", "生态成熟，文档完整", "约 100—300 元/年", "备选"),
    ("香港轻量服务器 + Nginx", "通常不需要备案", "约 100—300 元/年", "暂时不想备案"),
    ("国内轻量服务器 + Nginx", "可运行后端接口", "约 200—500 元/年", "未来需要 Flask 后端"),
]:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
doc.add_heading("三、实施步骤", level=1)
numbers([
    "购买独立域名并完成实名认证，建议 .com 或 .cn。",
    "在腾讯云或阿里云提交 ICP 备案，备案期间可继续使用 Vercel。",
    "创建 COS/OSS 存储桶，开启静态网站功能，首页设为 index.html，错误页设为 404.html。",
    "运行 build-site.cmd，生成 dist 目录。",
    "上传 dist 中的全部文件到国内存储桶。",
    "绑定域名，配置免费 HTTPS 证书和 CDN 加速。",
    "将 astro.config.mjs 和 src/data/site.ts 中的网址改为国内域名。",
    "重新构建、上传，并用多地测速工具验证访问。",
])
doc.add_heading("四、建议的国内版结构", level=1)
bullets([
    "首页：个人身份、核心项目、成果证据和国内联系方式。",
    "个人介绍：教育背景、综合测评、专业方向和能力结构。",
    "项目经历：堆叠卡片，突出个人分工和量化结果。",
    "作品展示：七个核心项目，详情页可放代码、报告和演示截图。",
    "数据图表：问卷、模型、仿真、性能等可核验数据。",
    "荣誉与实践：奖项、志愿服务、社团与 KAB 经历。",
])
doc.add_heading("五、域名与内容建议", level=1)
bullets([
    "建议使用真实姓名拼音域名，例如 xuchenyang.com。",
    "简历和 PPT 优先放国内域名，Vercel 地址作为备用地址。",
    "国内备案信息需填写真实主体和网站用途。",
    "网站公开内容不要放手机号、学号、出生年月和家庭住址。",
    "页面底部增加备案号，并链接工信部备案系统。",
])
doc.add_heading("六、预计时间和成本", level=1)
bullets([
    "域名：约 30—100 元/年。",
    "COS/OSS + CDN：低流量个人站通常约 100—300 元/年。",
    "域名实名：通常 1—3 天。",
    "ICP 备案：通常 7—20 个工作日。",
    "部署和测速：约半天。",
])
doc.add_heading("七、最推荐路线", level=1)
doc.add_paragraph("腾讯云域名 + ICP 备案 + COS 静态托管 + CDN + 免费 SSL。保留现有 Astro 代码，不重写网站；只增加国内部署目标和域名配置。")
doc.save(out)
print(out)
