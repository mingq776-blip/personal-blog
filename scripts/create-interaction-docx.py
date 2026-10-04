from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
out = Path(r"E:/codex/个人博客/docs/交互与版式调整说明.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 26, "145A2A"), ("Heading 1", 18, "237A3A")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph("交互与版式调整说明", style="Title")
doc.add_heading("一、绿色图片重新使用", level=1)
bullets([
    "森林图片作为全站固定模糊背景。",
    "湖泊图片保留在首页主视觉后方。",
    "草地图片作为项目经历堆叠卡片的背景，并增加浅绿色遮盖层。",
    "首页“自然的绿色片段”图片栏目已删除。",
])
doc.add_heading("二、项目经历改为可交互堆叠", level=1)
bullets([
    "6 段项目经历改为堆叠卡片，当前经历位于最上层。",
    "支持“下一项”“上一项”、圆点和点击当前卡片切换。",
    "切换时包含层叠、缩放、透明度和位置过渡。",
    "手机端保留相同交互，并增加卡片高度以保证内容完整。",
])
doc.add_heading("三、荣誉与实践弹窗", level=1)
bullets([
    "首页“荣誉与实践”卡片可以点击。",
    "点击后打开弹窗，完整展示 10 项奖学金、竞赛、志愿服务和资格证书。",
    "弹窗支持关闭按钮、点击空白区域关闭和键盘交互。",
])
doc.add_heading("四、作品展示取消图片", level=1)
bullets([
    "首页“03 作品展示”已移除项目封面图片。",
    "作品卡片改为编号、标题、量化指标、技术栈和证据口径组成的文字卡片。",
    "作品详情页仍保留封面图片，方便查看完整项目。",
])
doc.add_heading("五、可访问性", level=1)
bullets([
    "堆叠切换按钮和圆点已扩大点击区域。",
    "键盘焦点、弹窗关闭和减少动画兼容性已保留。",
    "首页 Lighthouse 无障碍评分 100。",
])
doc.add_heading("六、预览", level=1)
for image, caption in [
    (r"E:/codex/个人博客/docs/preview-home-interactive.png", "首页交互版预览"),
    (r"E:/codex/个人博客/docs/preview-awards-dialog.png", "荣誉与奖项弹窗预览"),
]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(image, width=Inches(6.2))
    c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.save(out)
print(out)
