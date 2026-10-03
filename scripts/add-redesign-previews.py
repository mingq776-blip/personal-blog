from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
path = r'E:\codex\个人博客\docs\网站改版说明.docx'
doc = Document(path)
doc.add_heading('六、设计预览', level=1)
for image, caption in [
    (r'E:\codex\个人博客\docs\preview-home-v2.png', '新版首页：杂志式封面、粒子动画与内容分区'),
    (r'E:\codex\个人博客\docs\preview-dashboard-v2.png', '数据图表页：趋势、能力分布与时间分配'),
]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(image, width=Inches(6.2))
    c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.save(path)
print(path)
