from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out = Path(r"E:/codex/个人博客/docs/首页标题与英文名称修改说明.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 26, "145A2A"), ("Heading 1", 18, "237A3A")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
doc.add_paragraph("首页标题与英文名称修改说明", style="Title")
doc.add_heading("一、修改内容", level=1)
doc.add_paragraph("1. “让数据变成作品”已改为可交互标题。")
doc.add_paragraph("鼠标悬浮时，7 个汉字会依次产生上浮、旋转和回弹动画，第二行同时增加绿色光晕。")
doc.add_paragraph("2. “个人作品档案”已替换为英文：")
doc.add_paragraph("Xu Chenyang | Personal Portfolio")
doc.add_paragraph("该英文名称同步用于浏览器标题和顶部品牌区。")
doc.add_heading("二、实现方式", level=1)
doc.add_paragraph("标题拆分为单个汉字元素，通过 CSS 动画和 --i 延迟变量形成逐字波动效果；动画不依赖 JavaScript，减少性能消耗。")
doc.add_heading("三、预览与部署", level=1)
doc.add_paragraph("本地预览：http://127.0.0.1:4321/")
doc.add_paragraph("最新代码同时提交到 GitHub，并由 Vercel 自动重新部署。")
doc.save(out)
print(out)
