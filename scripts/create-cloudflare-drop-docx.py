from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
out = Path(r"E:/codex/个人博客/docs/Cloudflare Drop使用说明.docx")
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.4
for name, size, color in [("Title", 26, "145A2A"), ("Heading 1", 18, "237A3A")]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
def numbers(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")
doc.add_paragraph("Cloudflare Drop 使用说明", style="Title")
doc.add_heading("一、结论", level=1)
doc.add_paragraph("适合做 Demo、临时预览、设计原型和快速分享；不适合作为中国大陆长期稳定展示的唯一方案。")
bullets([
    "无需 Cloudflare 账号即可先上传。",
    "支持直接拖入静态站点文件夹或 ZIP。",
    "部署后获得 workers.dev 临时地址。",
    "临时部署保留约 60 分钟，需在期限内登录并 Claim。",
])
doc.add_heading("二、本项目如何使用", level=1)
bullets([
    "已准备好上传包：E:\\codex\\个人博客\\outputs\\网站静态版-CloudflareDrop.zip",
    "打开网址：https://www.cloudflare.com/drop/",
    "把 ZIP 文件拖入页面并等待上传完成。",
    "复制生成的 workers.dev 预览地址。",
    "如果要保留，在 60 分钟内点击 Claim，登录 Cloudflare 完成认领。",
])
doc.add_heading("三、限制", level=1)
bullets([
    "只能部署纯静态 HTML、CSS、JS、图片和搜索索引。",
    "不能运行 Flask、Python、数据库或长期后台服务。",
    "青穗项目的 9 个 API 需要单独部署后端，Cloudflare Drop 不能直接运行 app.py。",
    "workers.dev 在中国大陆的访问稳定性无法保证。",
    "长期国内展示仍建议独立域名 + ICP 备案 + 国内 COS/OSS/CDN，或香港主机。",
])
doc.add_heading("四、CLI 临时部署示例", level=1)
p = doc.add_paragraph(); r = p.add_run('npm exec --yes wrangler@latest -- deploy ./dist --name xuchenyang-portfolio --temporary --compatibility-date 2026-10-05'); r.font.name = "Consolas"; r.font.size = Pt(8.5)
doc.add_heading("五、推荐用途", level=1)
bullets([
    "临时给别人发一个网页预览。",
    "答辩前做备用演示地址。",
    "测试新版页面是否能正常打开。",
    "正式长期国内展示仍使用国内云静态托管方案。",
])
doc.save(out)
print(out)
